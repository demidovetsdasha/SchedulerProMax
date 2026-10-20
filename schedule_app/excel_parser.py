"""Чтение расписания из Excel (.xlsx / .xls) с сохранением геометрии объединений.

Ожидаемая структура листа:
    строка с «День недели» | «Пара» | «Время пары» | 201/1 | 201/2 | 202/1 ...
    далее строки пар; день недели — объединённая ячейка на все строки дня.

Геометрия сохраняется полностью:
  * строки Excel группируются в «блоки пары» по объединённой ячейке времени
    (например, C10:C12 — одна пара с тремя дисциплинами по выбору в трёх строках);
  * для каждой ячейки групп запоминается её объединённый диапазон Excel, чтобы
    в Word воссоздать те же объединения по горизонтали и вертикали.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

DAY_NAMES = ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье"]

MergeRange = tuple[int, int, int, int]   # (r1, c1, r2, c2), 0-based, включительно


@dataclass
class Column:
    index: int          # индекс столбца на листе (0-based)
    name: str           # "201/1" — как в Excel
    group: str          # "201"


@dataclass
class Lesson:
    text: str
    merge: MergeRange | None   # диапазон объединения Excel, если ячейка объединена


@dataclass
class Block:
    """Одна пара: строки Excel под одной (возможно объединённой) ячейкой времени."""
    pair: str
    time: str
    rows: list[int] = field(default_factory=list)


@dataclass
class Day:
    name: str
    blocks: list[Block] = field(default_factory=list)


@dataclass
class SheetSchedule:
    sheet_name: str
    title: str
    columns: list[Column]
    days: list[Day]
    cells: dict[tuple[int, int], Lesson] = field(default_factory=dict)

    def lesson(self, row: int, col: int) -> Lesson | None:
        return self.cells.get((row, col))

    @property
    def groups(self) -> list[str]:
        seen: list[str] = []
        for c in self.columns:
            if c.group not in seen:
                seen.append(c.group)
        return seen

    @property
    def course(self) -> int | None:
        m = re.search(r"(\d)\s*-?\s*(?:й|ой|ый)?\s*курс", self.title, re.I)
        if m:
            return int(m.group(1))
        firsts = {g[0] for g in self.groups if g[:1].isdigit()}
        return int(firsts.pop()) if len(firsts) == 1 else None


@dataclass
class ExcelSchedule:
    path: Path
    sheets: list[SheetSchedule]


# --------------------------------------------------------------------------- utils

def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, dt.time):
        value = value.strftime("%H:%M")
    text = str(value).replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return "\n".join(line.strip() for line in text.splitlines()).strip()


def _normalize_day(text: str) -> str | None:
    t = text.strip().lower()
    for name in DAY_NAMES:
        if t.startswith(name.lower()[:3]):
            return name
    return None


# ------------------------------------------------------------------ sheet readers

class _Grid:
    """Унифицированное представление листа: значения + объединённые диапазоны."""

    def __init__(self, name: str, values: list[list], merges: list[MergeRange]):
        self.name = name
        self.values = values
        self.owner: dict[tuple[int, int], MergeRange] = {}
        for rng in merges:
            r1, c1, r2, c2 = rng
            if r1 == r2 and c1 == c2:
                continue
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    self.owner[(r, c)] = rng

    @property
    def nrows(self) -> int:
        return len(self.values)

    def ncols(self) -> int:
        return max((len(r) for r in self.values), default=0)

    def raw(self, r: int, c: int):
        row = self.values[r] if r < len(self.values) else []
        return row[c] if c < len(row) else None

    def value(self, r: int, c: int) -> str:
        rng = self.owner.get((r, c))
        if rng:
            return _clean(self.raw(rng[0], rng[1]))
        return _clean(self.raw(r, c))

    def merge(self, r: int, c: int) -> MergeRange | None:
        return self.owner.get((r, c))


def _read_xlsx(path: Path) -> list[_Grid]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    grids = []
    for ws in wb.worksheets:
        values = [list(row) for row in ws.iter_rows(values_only=True)]
        merges = [
            (m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1)
            for m in ws.merged_cells.ranges
        ]
        grids.append(_Grid(ws.title, values, merges))
    return grids


def _read_xls(path: Path) -> list[_Grid]:
    import xlrd

    book = xlrd.open_workbook(str(path), formatting_info=True)
    grids = []
    for sh in book.sheets():
        values = [sh.row_values(r) for r in range(sh.nrows)]
        merges = [(r1, c1, r2 - 1, c2 - 1) for r1, r2, c1, c2 in sh.merged_cells]
        grids.append(_Grid(sh.name, values, merges))
    return grids


# ------------------------------------------------------------------------ parsing

def _find_header(grid: _Grid) -> tuple[int, int] | None:
    for r in range(min(grid.nrows, 30)):
        for c in range(min(grid.ncols(), 5)):
            if grid.value(r, c).lower().startswith("день недели"):
                return r, c
    return None


def _parse_grid(grid: _Grid) -> SheetSchedule | None:
    found = _find_header(grid)
    if found is None:
        return None
    header_row, day_col = found

    # Колонки «Пара» и «Время пары» ищем по заголовку, остальные — подгруппы.
    pair_col = time_col = None
    columns: list[Column] = []
    for c in range(day_col + 1, grid.ncols()):
        head = grid.value(header_row, c)
        if not head:
            continue
        low = head.lower()
        if low.startswith("пара") and pair_col is None:
            pair_col = c
        elif low.startswith("время") and time_col is None:
            time_col = c
        else:
            name = head.replace(" ", "")
            columns.append(Column(index=c, name=name, group=name.split("/")[0]))

    title = ""
    for r in range(header_row):
        for c in range(grid.ncols()):
            v = grid.value(r, c)
            if v and "курс" in v.lower():
                title = v
                break
        if title:
            break
    if not title and header_row > 0:
        title = grid.value(0, 0)

    def block_key(r: int):
        """Строки одной пары имеют общий диапазон объединения ячейки времени (или пары)."""
        for c in (time_col, pair_col):
            if c is not None and grid.merge(r, c):
                return ("m", grid.merge(r, c)[0], c)
        return None

    days: list[Day] = []
    current: Day | None = None
    block: Block | None = None
    for r in range(header_row + 1, grid.nrows):
        day_text = grid.value(r, day_col)
        day_name = _normalize_day(day_text) if day_text else None
        if day_name and (current is None or current.name != day_name):
            current = Day(day_name)
            days.append(current)
            block = None
        if current is None:
            continue

        time = grid.value(r, time_col) if time_col is not None else ""
        pair = grid.value(r, pair_col) if pair_col is not None else ""
        key = block_key(r)
        same_block = block is not None and (
            (key is not None and key == block_key(block.rows[0]))
            # строка без своего времени/номера пары — продолжение предыдущей пары
            or (key is None and not time and not pair)
        )
        if same_block:
            block.rows.append(r)
        elif time or pair:
            block = Block(pair=pair, time=time, rows=[r])
            current.blocks.append(block)
        else:
            continue

    sheet = SheetSchedule(grid.name, title, columns, days)
    for day in days:
        for blk in day.blocks:
            for r in blk.rows:
                for col in columns:
                    text = grid.value(r, col.index)
                    rng = grid.merge(r, col.index)
                    if text or rng:
                        sheet.cells[(r, col.index)] = Lesson(text, rng)
    return sheet


def load_excel(path: str | Path) -> ExcelSchedule:
    path = Path(path)
    grids = _read_xls(path) if path.suffix.lower() == ".xls" else _read_xlsx(path)
    sheets = [s for s in (_parse_grid(g) for g in grids) if s and s.columns]
    if not sheets:
        raise ValueError("В файле не найден лист с расписанием (нет строки «День недели»).")
    return ExcelSchedule(path, sheets)
