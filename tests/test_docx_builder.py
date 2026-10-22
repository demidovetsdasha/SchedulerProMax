"""Разбиение на страницы и сборка итогового документа Word по примерам."""
from __future__ import annotations

import datetime as dt

import pytest
from docx import Document

from conftest import sample
from schedule_app.docx_builder import DocumentParams, academic_year, build_document, paginate
from schedule_app.excel_parser import Column


def _cols(*names: str) -> list[Column]:
    return [Column(index=i, name=n, group=n.split("/")[0]) for i, n in enumerate(names)]


def _pages(pages):
    return [[c.name for c in page] for page in pages]


# ------------------------------------------------------------------ paginate

def test_paginate_does_not_split_group():
    cols = _cols("201/1", "201/2", "202/1", "202/2", "202/3")
    assert _pages(paginate(cols, 4)) == [["201/1", "201/2"], ["202/1", "202/2", "202/3"]]


def test_paginate_fills_page():
    cols = _cols("201/1", "201/2", "202/1", "202/2")
    assert _pages(paginate(cols, 4)) == [["201/1", "201/2", "202/1", "202/2"]]


def test_paginate_non_numeric_group():
    cols = _cols("ТянУИЯ", "201/1", "201/2", "202/1")
    assert _pages(paginate(cols, 3)) == [["ТянУИЯ", "201/1", "201/2"], ["202/1"]]


def test_paginate_group_larger_than_page_stays_whole():
    cols = _cols("201/1", "201/2", "201/3")
    assert _pages(paginate(cols, 2)) == [["201/1", "201/2", "201/3"]]


def test_paginate_empty():
    assert paginate([], 8) == []


@pytest.mark.parametrize("day, year", [
    (dt.date(2026, 8, 31), "2026-2027"),
    (dt.date(2026, 12, 31), "2026-2027"),
    (dt.date(2027, 1, 10), "2026-2027"),
    (dt.date(2027, 7, 31), "2026-2027"),
])
def test_academic_year(day, year):
    assert academic_year(day) == year


# ----------------------------------------------------------- build_document

def _params(cfg, sheet, code):
    fac = cfg.faculty_by_code(code)
    start = dt.date(2026, 10, 26)
    return DocumentParams(
        course=sheet.course,
        faculty_header=fac.header_name,
        date_start=start,
        date_end=start + dt.timedelta(days=6),
        head_ckop_title=cfg.head_ckop_title,
        head_ckop_name=cfg.head_ckop_name,
        deputy_dean=fac.deputy_dean,
        max_columns_per_page=cfg.max_columns_per_page,
        approval_year=2026,
        table_font_max=cfg.table_font_max,
        table_font_min=cfg.table_font_min,
    )


def _all_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


@pytest.mark.parametrize("key, code", [
    ("ФКЯиК", "ФКЯиК"),
    ("ФРЯ", "ФРЯ"),
    ("перев", "ПФ"),
    ("ф_т_ан_яз,", "ФАЯ"),
    ("нем_яз", "ФНЯ"),
    ("зимний", "ФАЯ"),
])
def test_build_document_from_sample(cfg, load_sheet, tmp_path, key, code):
    sheet = load_sheet(sample(key))
    cols = sheet.columns_for_course(sheet.course)
    pages = [[c.name for c in p] for p in paginate(cols, cfg.max_columns_per_page)]
    out = tmp_path / "out.docx"

    sizes = build_document(cfg.template, sheet, pages, _params(cfg, sheet, code), out)

    assert out.exists()
    assert len(sizes) == len(pages)
    assert all(cfg.table_font_min <= s <= cfg.table_font_max for s in sizes)
    text = _all_text(Document(str(out)))
    for name in (c.name for c in cols):
        assert name in text, f"в документе нет столбца {name}"
    assert cfg.faculty_by_code(code).header_name in text


def test_build_document_includes_non_numeric_group(cfg, load_sheet, tmp_path):
    sheet = load_sheet(sample("ФКЯиК"))
    out = tmp_path / "tyan.docx"
    build_document(cfg.template, sheet, [["ТянУИЯ", "201/1", "201/2"]],
                   _params(cfg, sheet, "ФКЯиК"), out)
    text = _all_text(Document(str(out)))
    assert "ТянУИЯ" in text
    assert "202/1" not in text


def test_build_document_page_order_follows_sheet(cfg, load_sheet, tmp_path):
    sheet = load_sheet(sample("ФКЯиК"))
    out = tmp_path / "order.docx"
    build_document(cfg.template, sheet, [["201/1", "ТянУИЯ"]], _params(cfg, sheet, "ФКЯиК"), out)
    text = _all_text(Document(str(out)))
    assert text.index("ТянУИЯ") < text.index("201/1")


def test_build_document_skips_empty_pages(cfg, load_sheet, tmp_path):
    sheet = load_sheet(sample("ФРЯ"))
    sizes = build_document(cfg.template, sheet, [[], ["501/1"], []],
                           _params(cfg, sheet, "ФРЯ"), tmp_path / "one.docx")
    assert len(sizes) == 1


def test_build_document_no_groups_raises(cfg, load_sheet, tmp_path):
    sheet = load_sheet(sample("ФРЯ"))
    with pytest.raises(ValueError, match="не выбраны группы"):
        build_document(cfg.template, sheet, [[]], _params(cfg, sheet, "ФРЯ"), tmp_path / "x.docx")
